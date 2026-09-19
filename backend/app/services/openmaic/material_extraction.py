"""Material intake policy and text extraction.

This module is the only place that turns an uploaded file into the text a stage
may later cite, so it carries two rules that the rest of the slice depends on:

1. **The extension decides the format, not the client.** A browser that labels a
   PDF as ``text/plain`` must not be able to talk this layer into decoding binary
   bytes as UTF-8 and presenting the result as document content. The media type
   stored on the material is therefore derived here from the filename.
2. **Nothing is fabricated.** A format this deployment cannot parse is reported
   as ``unsupported`` with *no* text, never as ``extracted`` with a guess. The
   service refuses a payload whose status and text disagree, so an honest
   ``unsupported`` is the only way to record "we could not read this".

Over-cap input is refused by name rather than truncated: a silently shortened
document would later be cited as if it were complete.
"""

from __future__ import annotations

import hashlib
from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from typing import Final, Optional

from ..openmaic.fusion_errors import FusionDocumentRejected, FusionInvalidRequest

#: Mirrors ``openmaic-service/src/material/repository.ts``. The two must agree:
#: the gateway refuses first so the student gets a CampusMate-shaped message, and
#: the service enforces the same number so a direct caller cannot bypass it.
MAX_MATERIAL_BYTES: Final = 2 * 1024 * 1024
MAX_MATERIAL_TEXT_BYTES: Final = 512 * 1024
MAX_FILENAME_LENGTH: Final = 255
MAX_REFERENCE_COUNT: Final = 50

#: The extraction statuses the managed service accepts.
EXTRACTION_STATUSES: Final = ("extracted", "unsupported", "empty")

_PATH_SEPARATORS = ("/", "\\")

#: Extension -> media type. Every extension we recognise is listed, including the
#: ones we cannot parse: a stored material should say what it *is* even when no
#: text could be read out of it.
_MEDIA_TYPES: Final[dict[str, str]] = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".markdown": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".ppt": "application/vnd.ms-powerpoint",
    ".doc": "application/msword",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xls": "application/vnd.ms-excel",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".gif": "image/gif",
    ".mp4": "video/mp4",
    ".mp3": "audio/mpeg",
    ".zip": "application/zip",
}

#: The subset this deployment can actually read text out of. Everything else in
#: `_MEDIA_TYPES` is retained as metadata with an honest `unsupported` status;
#: PPTX in particular needs a real OOXML reader this deployment does not ship.
_PARSEABLE_SUFFIXES: Final[frozenset[str]] = frozenset({".txt", ".md", ".markdown", ".pdf", ".docx"})

_TEXT_DECODINGS: Final = ("utf-8-sig", "utf-8", "gb18030")


@dataclass(frozen=True)
class ExtractedMaterial:
    """One intake decision: what the file is, and what text came out of it."""

    filename: str
    media_type: str
    byte_size: int
    sha256: str
    extraction_status: str
    text: str


def sanitize_filename(raw: Optional[str]) -> str:
    """A filename is a name, never a path.

    Taking the basename would quietly accept ``../../etc/passwd`` and store
    ``passwd``; refusing says the request was wrong, which is what it was.
    """
    filename = (raw or "").strip()
    if not filename:
        raise FusionInvalidRequest("文件名不能为空")
    if len(filename) > MAX_FILENAME_LENGTH:
        raise FusionInvalidRequest(f"文件名不能超过 {MAX_FILENAME_LENGTH} 个字符")
    if any(separator in filename for separator in _PATH_SEPARATORS):
        raise FusionInvalidRequest("文件名不能包含路径分隔符")
    if filename in (".", ".."):
        raise FusionInvalidRequest("文件名不合法")
    return filename


def media_type_for(filename: str) -> str:
    """Derive the media type from the filename. Unknown extensions stay generic."""
    suffix = PurePosixPath(filename).suffix.lower()
    return _MEDIA_TYPES.get(suffix, "application/octet-stream")


def _extension(filename: str) -> str:
    return PurePosixPath(filename).suffix.lower()


def _decode_text(content: bytes) -> Optional[str]:
    """Decode a text file, or return ``None`` when no known encoding fits.

    ``errors="strict"`` on purpose: a lossy decode would substitute replacement
    characters and hand back text the student never wrote.
    """
    for encoding in _TEXT_DECODINGS:
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return None


def _extract_pdf(content: bytes) -> Optional[str]:
    try:
        from PyPDF2 import PdfReader
    except ImportError:  # pragma: no cover - depends on deployment extras
        return None
    try:
        reader = PdfReader(BytesIO(content))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception:
        # A corrupt or encrypted PDF yields no text; that is "unsupported", not
        # an error the student can act on.
        return None
    text = "\n".join(pages).strip()
    return text or None


def _extract_docx(content: bytes) -> Optional[str]:
    try:
        from docx import Document
    except ImportError:  # pragma: no cover - depends on deployment extras
        return None
    try:
        document = Document(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                blocks.append("\t".join(cell.text for cell in row.cells))
    except Exception:
        return None
    text = "\n".join(blocks).strip()
    return text or None


def _normalize(text: str) -> str:
    return text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "").strip()


def extract_material(*, filename: str, content: bytes) -> ExtractedMaterial:
    """Turn one uploaded file into the record the managed service will store."""
    name = sanitize_filename(filename)
    byte_size = len(content)
    if byte_size > MAX_MATERIAL_BYTES:
        raise FusionDocumentRejected(
            details={"limit": "max_material_bytes", "max_bytes": MAX_MATERIAL_BYTES}
        )
    digest = hashlib.sha256(content).hexdigest()
    media_type = media_type_for(name)

    def build(status: str, text: str) -> ExtractedMaterial:
        if status == "extracted" and len(text.encode("utf-8")) > MAX_MATERIAL_TEXT_BYTES:
            # Refused by name rather than truncated: a shortened document would
            # later be cited as though it were whole.
            raise FusionDocumentRejected(
                details={"limit": "max_material_text_bytes", "max_bytes": MAX_MATERIAL_TEXT_BYTES}
            )
        return ExtractedMaterial(
            filename=name,
            media_type=media_type,
            byte_size=byte_size,
            sha256=digest,
            extraction_status=status,
            text=text if status == "extracted" else "",
        )

    if byte_size == 0:
        return build("empty", "")

    suffix = _extension(name)
    if suffix not in _PARSEABLE_SUFFIXES:
        # Known-unparseable and entirely unknown extensions take the same branch:
        # the file is kept as metadata and the student is told it has no readable
        # text. Media type still reflects what the file is, when we know.
        return build("unsupported", "")
    if media_type.startswith("text/"):
        decoded = _decode_text(content)
        if decoded is None:
            return build("unsupported", "")
        text = _normalize(decoded)
        return build("extracted" if text else "empty", text)
    if suffix == ".pdf":
        text = _extract_pdf(content)
        return build("extracted", _normalize(text)) if text else build("unsupported", "")
    if suffix == ".docx":
        text = _extract_docx(content)
        return build("extracted", _normalize(text)) if text else build("unsupported", "")
    return build("unsupported", "")


def is_supported_extension(filename: str) -> bool:
    """True when this deployment may be able to read text out of the file."""
    return _extension(filename) in _PARSEABLE_SUFFIXES


__all__ = [
    "EXTRACTION_STATUSES",
    "ExtractedMaterial",
    "MAX_FILENAME_LENGTH",
    "MAX_MATERIAL_BYTES",
    "MAX_MATERIAL_TEXT_BYTES",
    "MAX_REFERENCE_COUNT",
    "extract_material",
    "is_supported_extension",
    "media_type_for",
    "sanitize_filename",
]
